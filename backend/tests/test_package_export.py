"""Phase 9-C: material package volume export (成卷输出).

Covers the backend acceptance list from the Phase 9-C design (§8.1):
snapshot binding & selection, explicit precondition errors, original
content embedding with visible failure markers, the neutral A3 appendix,
ownership, listing and the export.created audit row.
"""

from __future__ import annotations

import io
import uuid

from tests.helpers import make_txt, upload
from tests.test_material_packages import _mk_package, _mk_evidence, _save, _snapshot


def _err_code(resp) -> str:
    return resp.json()["error"]["code"]


def _file_backed_package(client, headers, case_id, name="仲裁材料包"):
    """One analysed evidence row + package + snapshot v1 (draft v2)."""
    data = make_txt("月工资 9000 元 加班费约定")
    fid = upload(client, headers, case_id, "劳动合同.txt", data).json()["files"][0]["id"]
    client.post(f"/api/v1/files/{fid}/analyze", headers=headers)
    rows = client.get(f"/api/v1/cases/{case_id}/evidence", headers=headers).json()["evidence"]
    for row in rows:
        client.post(f"/api/v1/evidence/{row['id']}/confirm", headers=headers)

    pkg = _mk_package(client, headers, case_id, name=name)
    saved = _save(
        client,
        headers,
        pkg["id"],
        pkg["draft_version"],
        [{"evidence_id": r["id"]} for r in rows],
    )
    assert saved.status_code == 200, saved.text
    snap = _snapshot(client, headers, pkg["id"], saved.json()["draft_version"])
    assert snap.status_code == 201, snap.text
    return pkg, rows[0]


def _create_export(client, headers, pkg_id, **body):
    return client.post(
        f"/api/v1/material-packages/{pkg_id}/exports", json=body, headers=headers
    )


def _ready_package(client, headers, case_id, name="仲裁材料包", evidence_title="工资条"):
    """Package with one item, snapshotted to v1 (empty packages cannot snapshot)."""
    pkg = _mk_package(client, headers, case_id, name=name)
    ev = _mk_evidence(client, headers, case_id, title=evidence_title)
    saved = _save(
        client, headers, pkg["id"], pkg["draft_version"], [{"evidence_id": ev}]
    )
    assert saved.status_code == 200, saved.text
    snap = _snapshot(client, headers, pkg["id"], saved.json()["draft_version"])
    assert snap.status_code == 201, snap.text
    return pkg


def _content_disposition_name(resp) -> str:
    """Decode the RFC 5987 UTF-8 filename from Content-Disposition."""
    import re
    from urllib.parse import unquote

    m = re.search(r"filename\*=UTF-8''(\S+)", resp.headers["Content-Disposition"])
    assert m, resp.headers["Content-Disposition"]
    return unquote(m.group(1))


def _docx_text(client, headers, job_id) -> str:
    data = client.get(f"/api/v1/exports/{job_id}/download", headers=headers).content
    import docx

    doc = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            text += "\n" + " ".join(c.text for c in row.cells)
    for section in doc.sections:
        text += "\n" + "\n".join(p.text for p in section.footer.paragraphs)
    return text


# ---------- preconditions & error envelope ----------


def test_create_export_requires_auth(client, case_id):
    # no auth at all — the security dependency must reject first
    resp = client.post(
        "/api/v1/material-packages/00000000-0000-0000-0000-000000000001/exports",
        json={"format": "pdf"},
    )
    assert resp.status_code == 401


def test_export_requires_snapshot_409(client, auth_headers, case_id):
    pkg = _mk_package(client, auth_headers, case_id)
    _save(
        client,
        auth_headers,
        pkg["id"],
        pkg["draft_version"],
        [{"evidence_id": _mk_evidence(client, auth_headers, case_id)}],
    )
    resp = _create_export(client, auth_headers, pkg["id"], format="pdf")
    assert resp.status_code == 409
    assert _err_code(resp) == "SNAPSHOT_REQUIRED"


def test_export_unknown_snapshot_version_404(client, auth_headers, case_id):
    pkg = _ready_package(client, auth_headers, case_id)
    resp = _create_export(client, auth_headers, pkg["id"], format="pdf", snapshot_version=99)
    assert resp.status_code == 404
    assert _err_code(resp) == "SNAPSHOT_NOT_FOUND"


def test_export_rejects_zip_format(client, auth_headers, case_id):
    pkg = _ready_package(client, auth_headers, case_id)
    resp = _create_export(client, auth_headers, pkg["id"], format="zip")
    assert resp.status_code == 422
    assert _err_code(resp) == "UNSUPPORTED_FORMAT"


def test_gap_appendix_unavailable_409(client, auth_headers, case_id):
    pkg = _ready_package(client, auth_headers, case_id)
    resp = _create_export(
        client, auth_headers, pkg["id"], format="pdf", include_gap_appendix=True
    )
    assert resp.status_code == 409
    assert _err_code(resp) == "GAP_APPENDIX_UNAVAILABLE"


def test_export_ownership_404(client, auth_headers, second_user_headers, case_id):
    pkg = _ready_package(client, auth_headers, case_id)
    resp = _create_export(client, second_user_headers, pkg["id"], format="pdf")
    assert resp.status_code == 404
    list_resp = client.get(
        f"/api/v1/material-packages/{pkg['id']}/exports",
        headers=second_user_headers,
    )
    assert list_resp.status_code == 404


# ---------- happy paths ----------


def test_pdf_export_completes_with_snapshot_binding(client, auth_headers, case_id):
    pkg, _ = _file_backed_package(client, auth_headers, case_id)
    resp = _create_export(client, auth_headers, pkg["id"], format="pdf")
    assert resp.status_code == 201, resp.text
    job = resp.json()
    assert job["status"] == "COMPLETED"
    assert job["completed_at"] is not None
    assert job["package_id"] == pkg["id"]
    assert job["snapshot_id"] is not None
    assert job["snapshot_version"] == 1
    assert job["include_gap_appendix"] is False

    dl = client.get(f"/api/v1/exports/{job['id']}/download", headers=auth_headers)
    assert dl.status_code == 200
    assert dl.content[:5] == b"%PDF-"
    assert "attachment" in dl.headers["Content-Disposition"]
    assert _content_disposition_name(dl).startswith("材料包-")

    # 内容抽查：正文含原始内容与卷面要素，不含内部审核状态（验收 §五-7/10）。
    import pymupdf

    with pymupdf.open(stream=dl.content, filetype="pdf") as pdf:
        text = "\n".join(page.get_text() for page in pdf)
    assert "材料包：仲裁材料包" in text
    assert "月工资 9000 元" in text
    assert "拟证明事项" in text
    assert "review_status" not in text
    assert "CONFIRMED" not in text


def test_pdf_footer_shows_true_total_page_count(client, auth_headers, case_id):
    """Regression: every footer must read 共 <total> 页, not the current page.

    A 2-item package renders as index page + one page per material = 3 pages.
    The old footer used ReportLab's `doc.page`, which during the build is only
    the *current* page number, so each page claimed 共 1 页 / 共 2 页 / 共 3 页.
    """
    pkg = _mk_package(client, auth_headers, case_id, name="页码卷")
    ev1 = _mk_evidence(client, auth_headers, case_id, title="工资条", proof="用于证明工资数额。")
    ev2 = _mk_evidence(client, auth_headers, case_id, title="考勤表", proof="用于证明出勤情况。")
    saved = _save(
        client,
        auth_headers,
        pkg["id"],
        pkg["draft_version"],
        [{"evidence_id": ev1}, {"evidence_id": ev2}],
    )
    assert saved.status_code == 200, saved.text
    snap = _snapshot(client, auth_headers, pkg["id"], saved.json()["draft_version"])
    assert snap.status_code == 201, snap.text

    resp = _create_export(client, auth_headers, pkg["id"], format="pdf")
    assert resp.status_code == 201, resp.text
    dl = client.get(f"/api/v1/exports/{resp.json()['id']}/download", headers=auth_headers)
    assert dl.status_code == 200

    import pymupdf

    with pymupdf.open(stream=dl.content, filetype="pdf") as pdf:
        assert pdf.page_count == 3, f"expected a 3-page volume, got {pdf.page_count}"
        page_texts = [page.get_text("text") for page in pdf]

    # Every page shows the real total…
    for i, text in enumerate(page_texts, start=1):
        assert "共 3 页" in text, f"page {i} footer missing 共 3 页: {text[-80:]!r}"
    # …and the correct current page.
    assert "第 1 页 / 共 3 页" in page_texts[0]
    assert "第 2 页 / 共 3 页" in page_texts[1]
    assert "第 3 页 / 共 3 页" in page_texts[2]


def test_docx_export_embeds_original_content_and_index(client, auth_headers, case_id):
    pkg, _ = _file_backed_package(client, auth_headers, case_id, name="考勤卷")
    resp = _create_export(client, auth_headers, pkg["id"], format="docx")
    assert resp.status_code == 201, resp.text

    text = _docx_text(client, auth_headers, resp.json()["id"])
    assert "材料包：考勤卷" in text
    assert "生成时间" in text
    # 卷面目录表
    assert "包内编号" in text
    assert "拟证明事项" in text
    assert "原证据" in text
    # per-item section + original content embedding
    assert "原始内容：" in text
    assert "月工资 9000 元" in text
    # A3 result data must never leak into the export (decision 4.1)
    assert "review_status" not in text
    # compliance footer
    assert "不构成正式法律意见" in text


def test_export_from_specific_snapshot_version(client, auth_headers, case_id):
    pkg, ev = _file_backed_package(client, auth_headers, case_id)
    saved = _save(
        client,
        auth_headers,
        pkg["id"],
        2,
        [{"evidence_id": ev["id"], "display_name": "修订后的材料名"}],
    )
    assert saved.status_code == 200, saved.text
    snap2 = _snapshot(client, auth_headers, pkg["id"], saved.json()["draft_version"])
    assert snap2.status_code == 201

    v1 = _create_export(
        client, auth_headers, pkg["id"], format="docx", snapshot_version=1
    )
    assert v1.status_code == 201
    assert v1.json()["snapshot_version"] == 1
    v1_text = _docx_text(client, auth_headers, v1.json()["id"])
    assert "修订后的材料名" not in v1_text

    latest = _create_export(client, auth_headers, pkg["id"], format="docx")
    assert latest.json()["snapshot_version"] == 2
    latest_text = _docx_text(client, auth_headers, latest.json()["id"])
    assert "修订后的材料名" in latest_text


def test_gap_appendix_rendered_neutrally(client, auth_headers, case_id):
    chat = _mk_evidence(
        client,
        auth_headers,
        case_id,
        title="聊天记录",
        proof="用于证明入职时间。",
    )
    # manual evidence carries the fact label in its summary
    client.patch(
        f"/api/v1/evidence/{chat}",
        json={"summary": "入职时间：2022-10-01"},
        headers=auth_headers,
    )
    gap = client.post(f"/api/v1/cases/{case_id}/evidence-gaps", headers=auth_headers)
    assert gap.status_code == 200, gap.text
    assert gap.json()["status"] == "SUCCESS"

    pkg = _mk_package(client, auth_headers, case_id)
    saved = _save(
        client, auth_headers, pkg["id"], pkg["draft_version"], [{"evidence_id": chat}]
    )
    assert saved.status_code == 200, saved.text
    snap = _snapshot(client, auth_headers, pkg["id"], saved.json()["draft_version"])
    assert snap.status_code == 201, snap.text
    resp = _create_export(
        client, auth_headers, pkg["id"], format="docx", include_gap_appendix=True
    )
    assert resp.status_code == 201, resp.text
    job = resp.json()
    assert job["include_gap_appendix"] is True
    assert job["gap_result_id"] is not None

    text = _docx_text(client, auth_headers, job["id"])
    assert "附录：材料完整性提示（系统生成）" in text
    assert "不构成事实认定" in text
    assert "review_status" not in text

    # without the toggle no appendix is rendered
    plain = _create_export(client, auth_headers, pkg["id"], format="docx")
    plain_text = _docx_text(client, auth_headers, plain.json()["id"])
    assert "附录：材料完整性提示" not in plain_text


def test_appendix_marked_stale_after_evidence_edit(client, auth_headers, case_id):
    ev_id = _mk_evidence(
        client, auth_headers, case_id, title="劳动合同", proof="证明劳动关系。"
    )
    client.patch(
        f"/api/v1/evidence/{ev_id}",
        json={"summary": "入职时间：2022-10-01"},
        headers=auth_headers,
    )
    gap = client.post(f"/api/v1/cases/{case_id}/evidence-gaps", headers=auth_headers)
    assert gap.status_code == 200

    client.patch(
        f"/api/v1/evidence/{ev_id}", json={"title": "劳动合同（已修改）"}, headers=auth_headers
    )

    pkg = _mk_package(client, auth_headers, case_id)
    saved = _save(
        client, auth_headers, pkg["id"], pkg["draft_version"], [{"evidence_id": ev_id}]
    )
    assert saved.status_code == 200, saved.text
    snap = _snapshot(client, auth_headers, pkg["id"], saved.json()["draft_version"])
    assert snap.status_code == 201, snap.text
    resp = _create_export(
        client, auth_headers, pkg["id"], format="docx", include_gap_appendix=True
    )
    assert resp.status_code == 201, resp.text
    text = _docx_text(client, auth_headers, resp.json()["id"])
    assert "可能过期" in text


def test_missing_original_file_marked_explicitly(client, auth_headers, case_id):
    pkg, _ = _file_backed_package(client, auth_headers, case_id)
    files = client.get(f"/api/v1/cases/{case_id}/files", headers=auth_headers)
    assert files.status_code == 200, files.text
    file_rows = files.json()["files"]
    assert file_rows
    resp = client.delete(f"/api/v1/files/{file_rows[0]['id']}", headers=auth_headers)
    assert resp.status_code == 204, resp.text

    job = _create_export(client, auth_headers, pkg["id"], format="docx")
    assert job.status_code == 201, job.text
    text = _docx_text(client, auth_headers, job.json()["id"])
    assert "原始文件已不存在" in text or "无法可靠转换为正文" in text


# ---------- listing & audit ----------


def test_list_package_exports(client, auth_headers, case_id):
    pkg, _ = _file_backed_package(client, auth_headers, case_id)
    other = _ready_package(client, auth_headers, case_id, name="另一个材料包")
    other_job = _create_export(client, auth_headers, other["id"], format="docx")
    assert other_job.status_code == 201, other_job.text

    first = _create_export(client, auth_headers, pkg["id"], format="pdf")
    second = _create_export(client, auth_headers, pkg["id"], format="docx")
    assert first.status_code == second.status_code == 201

    resp = client.get(f"/api/v1/material-packages/{pkg['id']}/exports", headers=auth_headers)
    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] == 2
    assert len(body["exports"]) == 2
    # newest first
    assert body["exports"][0]["id"] == second.json()["id"]
    assert body["exports"][1]["id"] == first.json()["id"]
    ids = {e["id"] for e in body["exports"]}
    assert other_job.json()["id"] not in ids
    for e in body["exports"]:
        assert e["package_id"] == pkg["id"]
        assert e["snapshot_version"] == 1


def test_export_created_is_audited(client, auth_headers, case_id, db_session):
    from app.models.audit import AuditLog

    pkg, _ = _file_backed_package(client, auth_headers, case_id)
    job = _create_export(client, auth_headers, pkg["id"], format="pdf")
    assert job.status_code == 201

    row = (
        db_session.query(AuditLog)
        .filter(
            AuditLog.action == "export.created",
            AuditLog.entity_id == uuid.UUID(job.json()["id"]),
        )
        .one()
    )
    assert row.metadata_json["package_id"] == pkg["id"]
    assert row.metadata_json["snapshot_version"] == 1


# ---------- download filename compatibility (9-C 收尾验收 §三) ----------


def _ascii_fallback_name(resp) -> str:
    import re

    m = re.search(r'filename="([^"]*)"', resp.headers["Content-Disposition"])
    assert m, resp.headers["Content-Disposition"]
    return m.group(1)


def test_download_filename_ascii_fallback_english_name(client, auth_headers, case_id):
    """英文+空格包名：ASCII 回退保留原名（不再带前导 `-`），头值可 latin-1 编码。"""
    pkg = _ready_package(client, auth_headers, case_id, name="March Bundle v2")
    job = _create_export(client, auth_headers, pkg["id"], format="docx")
    assert job.status_code == 201, job.text

    dl = client.get(f"/api/v1/exports/{job.json()['id']}/download", headers=auth_headers)
    assert dl.status_code == 200
    header = dl.headers["Content-Disposition"]
    header.encode("latin-1")  # HTTP 头值硬性要求
    assert _ascii_fallback_name(dl) == "March Bundle v2.docx"
    assert _content_disposition_name(dl) == "材料包-March Bundle v2.docx"


def test_download_filename_special_characters_neutralised(client, auth_headers, case_id):
    """分号/引号包名：回退中无 `"`、`;`、`/`；filename* 精确往返。"""
    pkg = _ready_package(client, auth_headers, case_id, name='Q1; "final" draft')
    job = _create_export(client, auth_headers, pkg["id"], format="pdf")
    assert job.status_code == 201, job.text

    dl = client.get(f"/api/v1/exports/{job.json()['id']}/download", headers=auth_headers)
    assert dl.status_code == 200
    fallback = _ascii_fallback_name(dl)
    assert '"' not in fallback and ";" not in fallback and "/" not in fallback
    assert fallback.isascii() and fallback.strip()
    # _safe_filename 只剥引号；分号保留并由 filename* 精确携带
    assert _content_disposition_name(dl) == "材料包-Q1; final draft.pdf"


# ---------- snapshot immutability on export (9-C 收尾验收 §五-6) ----------


def test_export_without_version_binds_latest_snapshot_not_draft(
    client, auth_headers, case_id
):
    """未传 snapshot_version → 绑定最新快照而非草稿；草稿继续修改不泄漏。

    系统不会自动保存/自动创建快照：无快照时 409 SNAPSHOT_REQUIRED（另有
    test_export_requires_snapshot_409），有快照时导出内容永远来自不可变快照。
    """
    pkg, ev = _file_backed_package(client, auth_headers, case_id)  # 快照 v1 / 草稿 v2
    saved = _save(
        client,
        auth_headers,
        pkg["id"],
        2,
        [{"evidence_id": ev["id"], "display_name": "草稿新名称（不应出现在导出中）"}],
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["draft_version"] == 3  # 草稿前移，但没有新快照

    job = _create_export(client, auth_headers, pkg["id"], format="docx")
    assert job.status_code == 201, job.text
    assert job.json()["snapshot_version"] == 1  # 绑定快照 v1，而非草稿 v3
    text = _docx_text(client, auth_headers, job.json()["id"])
    assert "草稿新名称" not in text


def test_package_render_failure_records_failed_job(client, auth_headers, case_id, monkeypatch):
    """渲染失败：返回 RENDER_FAILED 422，作业行以 FAILED + 错误信息留存。"""
    from app.services import export as export_service

    pkg, _ = _file_backed_package(client, auth_headers, case_id)

    def boom(data):
        raise RuntimeError("simulated renderer crash")

    monkeypatch.setattr(export_service, "render_package_docx", boom)
    resp = _create_export(client, auth_headers, pkg["id"], format="docx")
    assert resp.status_code == 422, resp.text
    assert _err_code(resp) == "RENDER_FAILED"

    listed = client.get(
        f"/api/v1/material-packages/{pkg['id']}/exports", headers=auth_headers
    )
    assert listed.json()["total"] == 1
    failed = listed.json()["exports"][0]
    assert failed["status"] == "FAILED"
    assert "simulated renderer crash" in (failed["error_message"] or "")
