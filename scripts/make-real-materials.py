#!/usr/bin/env python3
"""Generate the Phase 9-D desensitised material suite (M1-M12).

Produces realistic-format but fictional evidence files plus a manifest.json
describing each file's expected pipeline behaviour. Everything is generated
at runtime into an output directory — no binaries are committed.

Usage:
    python scripts/make-real-materials.py <outdir>

Self-checks after generation: magic bytes, minimum sizes, UTF-8 validity.
Re-running overwrites the previous output (idempotent).
"""

from __future__ import annotations

import hashlib
import io
import json
import sys
from pathlib import Path

import pymupdf
import docx
from PIL import Image, ImageDraw, ImageFont

FONT = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"

# ---- M1 / M11 / M12 shared building blocks ---------------------------------

M1_TEXT = """劳动合同
甲方：某某科技有限公司
乙方：张三
合同期限：2022年10月1日至2025年9月30日
入职日期：2022年10月1日
月工资7500元
工作内容：软件工程师
"""

M11_TEXT = M1_TEXT + "附页：试用期考核合格，转正日期2023年1月1日。\n"

M12_PAGE1 = """第一页：劳动合同首页
甲方：某某科技有限公司
入职日期：2022年10月1日
月工资7500元
"""

M12_PAGE3 = """第三页：补充条款
甲方：某某科技有限公司
合同期限顺延至2026年9月30日。
"""

M2_TEXT = """工资条
员工：张三
所属期间：2024年3月
基本工资：6000元
绩效奖金：1500元
加班费：800元
应发合计：8300元
实发工资：8100元
"""

M3_TEXT = """微信聊天记录（导出）
[2024/3/15 10:23] 张三：领导，这个月工资什么时候发？
[2024/3/15 10:25] 李四：再等等，公司最近困难。
[2024/3/16 09:12] 张三：那社保还正常缴吗？
[2024/3/16 09:15] 李四：社保没问题，这个你放心。
[2024/3/18 17:40] 张三：如果一直发不出工资，我可能要申请仲裁。
[2024/3/18 17:42] 李四：别急，我们尽量协调。
[2024/3/20 11:00] 张三：好的，谢谢。
"""

M12_IMAGE_MARKER = "第二页扫描专用文字"


# ---- helpers ---------------------------------------------------------------


def _pdf_with_text(text: str) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_font(fontname="F0", fontfile=FONT)
    page.insert_text((50, 72), text, fontsize=10, fontname="F0")
    doc.subset_fonts()
    data = doc.tobytes()
    doc.close()
    return data


def _pdf_with_image_only(png: bytes) -> bytes:
    """A page whose only content is a raster image — no text layer at all."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(pymupdf.Rect(40, 40, 555, 800), stream=png)
    data = doc.tobytes()
    doc.close()
    return data


def _pdf_mixed(png: bytes) -> bytes:
    """3 pages: text / image-only / text — exercises per-page OCR fallback."""
    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_font(fontname="F0", fontfile=FONT)
    p1.insert_text((50, 72), M12_PAGE1, fontsize=10, fontname="F0")
    p2 = doc.new_page(width=595, height=842)
    p2.insert_image(pymupdf.Rect(40, 40, 555, 800), stream=png)
    p3 = doc.new_page()
    p3.insert_font(fontname="F0", fontfile=FONT)
    p3.insert_text((50, 72), M12_PAGE3, fontsize=10, fontname="F0")
    doc.subset_fonts()
    data = doc.tobytes()
    doc.close()
    return data


def _docx_attendance() -> bytes:
    d = docx.Document()
    d.add_paragraph("考勤记录")
    d.add_paragraph("员工：张三")
    d.add_paragraph("所属期间：2024年3月")
    table = d.add_table(rows=6, cols=4)
    rows = [
        ("日期", "上班打卡", "下班打卡", "状态"),
        ("2024-03-04", "09:02", "18:01", "正常"),
        ("2024-03-05", "09:42", "18:00", "迟到"),
        ("2024-03-06", "09:00", "20:30", "加班"),
        ("2024-03-07", "—", "—", "请假"),
        ("2024-03-08", "08:58", "18:05", "正常"),
    ]
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.rows[r].cells[c].text = value
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT, size)


def _scanned_page_png() -> bytes:
    """A 'scanned' page whose only text lives inside raster pixels.

    The marker string is drawn with a CJK font so it is visually present;
    mock OCR returns empty, so pipeline raw_text must NOT contain it.
    """
    img = Image.new("RGB", (595, 842), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.rectangle([40, 40, 555, 800], outline=(60, 60, 60), width=2)
    draw.text((80, 100), M12_IMAGE_MARKER, font=_font(32), fill=(0, 0, 0))
    draw.text((80, 200), "（扫描件页面，无文本层）", font=_font(20), fill=(80, 80, 80))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _chat_screenshot_png() -> bytes:
    """Fake chat UI screenshot: bubbles + timestamps, drawn with CJK font."""
    img = Image.new("RGB", (750, 900), (245, 245, 245))
    draw = ImageDraw.Draw(img)
    draw.rectangle([0, 0, 750, 70], fill=(50, 50, 50))
    draw.text((260, 20), "微信聊天", font=_font(26), fill=(255, 255, 255))
    bubbles = [
        (40, 110, "对方", "#e8e8e8", "这个月工资什么时候发？"),
        (40, 250, "对方", "#e8e8e8", "社保还正常缴吗？"),
        (390, 180, "我", "#95ec69", "再等等，公司最近困难。"),
        (390, 320, "我", "#95ec69", "社保没问题。"),
        (40, 460, "对方", "#e8e8e8", "一直发不出的话我要申请仲裁。"),
        (390, 600, "我", "#95ec69", "别急，尽量协调。"),
        (40, 740, "对方", "#e8e8e8", "好的，谢谢。"),
    ]
    for x, y, who, color, text in bubbles:
        draw.text((x, y - 28), who, font=_font(16), fill=(120, 120, 120))
        draw.rounded_rectangle([x, y, x + 320, y + 80], radius=12, fill=color)
        draw.text((x + 16, y + 24), text, font=_font(18), fill=(20, 20, 20))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _wage_photo_jpeg() -> bytes:
    """Fake wage-slip phone photo: white paper, ruled lines, amounts."""
    img = Image.new("RGB", (800, 1100), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.rectangle([30, 30, 770, 1070], outline=(80, 80, 80), width=3)
    draw.text((280, 70), "工 资 条", font=_font(40), fill=(0, 0, 0))
    lines = [
        "员工：张三",
        "所属期间：2024年3月",
        "基本工资：6000元",
        "绩效奖金：1500元",
        "加班费：800元",
        "应发合计：8300元",
        "实发工资：8100元",
    ]
    for i, line in enumerate(lines):
        draw.text((120, 180 + i * 110), line, font=_font(28), fill=(30, 30, 30))
        draw.line([110, 230 + i * 110, 690, 230 + i * 110], fill=(180, 180, 180), width=1)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    return buf.getvalue()


def _magic_ok(name: str, data: bytes) -> bool:
    if name.endswith(".pdf"):
        return data[:5] == b"%PDF-"
    if name.endswith(".png"):
        return data[:8] == b"\x89PNG\r\n\x1a\n"
    if name.endswith((".jpg", ".jpeg")):
        return data[:3] == b"\xff\xd8\xff"
    if name.endswith(".docx"):
        return data[:4] == b"PK\x03\x04"
    if name.endswith(".txt"):
        try:
            data.decode("utf-8")
            return True
        except UnicodeDecodeError:
            return False
    return False


# ---- material definitions --------------------------------------------------


def generate(outdir: Path) -> list[dict]:
    outdir.mkdir(parents=True, exist_ok=True)
    chat_png = _chat_screenshot_png()
    materials: list[tuple[str, str, bytes, dict]] = []

    materials.append(
        (
            "M1",
            "劳动合同.pdf",
            _pdf_with_text(M1_TEXT),
            {
                "kind": "good",
                "upload": "created",
                "analyze": "SUCCESS",
                "classification": "labor_contract",
                "confidence_min": 0.80,
                "dates_include": ["2022-10-01"],
                "amounts_include": ["7500元"],
                "orgs_include": ["某某科技有限公司"],
                "raw_text_include": ["入职日期：2022年10月1日"],
                "note": "带文本层 PDF；主链路材料",
            },
        )
    )
    materials.append(
        (
            "M2",
            "工资条.txt",
            M2_TEXT.encode("utf-8"),
            {
                "kind": "good",
                "upload": "created",
                "analyze": "SUCCESS",
                "classification": "salary_statement",
                "amounts_min": 3,
                "note": "纯文本工资条，金额抽取>=3",
            },
        )
    )
    materials.append(
        (
            "M3",
            "聊天记录.txt",
            M3_TEXT.encode("utf-8"),
            {
                "kind": "good",
                "upload": "created",
                "analyze": "SUCCESS",
                "classification": "chat_record",
                "confidence_min": 0.80,
                "note": "多行时间戳结构→chat_record（结构分3→置信0.96，非低置信）",
            },
        )
    )
    materials.append(
        (
            "M4",
            "考勤表.docx",
            _docx_attendance(),
            {
                "kind": "good",
                "upload": "created",
                "analyze": "SUCCESS",
                "classification": "attendance_record",
                "raw_text_include": ["09:42", "迟到"],
                "note": "DOCX 表格单元格进入 extracted_text",
            },
        )
    )
    materials.append(
        (
            "M5",
            "聊天截图.png",
            chat_png,
            {
                "kind": "good",
                "upload": "created",
                "analyze": "SUCCESS",
                "classification": "other",
                "confidence": 0.3,
                "raw_text_empty": True,
                "evidence_category": "other",
                "evidence_summary": "材料内容为空或无法识别。",
                "note": "mock OCR 恒空→SUCCESS+空文本+other/0.3",
            },
        )
    )
    materials.append(
        (
            "M6",
            "工资条照片.jpg",
            _wage_photo_jpeg(),
            {
                "kind": "good",
                "upload": "created",
                "analyze": "SUCCESS",
                "classification": "other",
                "confidence": 0.3,
                "raw_text_empty": True,
                "evidence_category": "other",
                "note": "同 M5（JPEG 路径）",
            },
        )
    )
    materials.append(
        (
            "M7",
            "扫描件.pdf",
            _pdf_with_image_only(chat_png),
            {
                "kind": "good",
                "upload": "created",
                "analyze": "SUCCESS",
                "classification": "other",
                "confidence": 0.3,
                "raw_text_empty": True,
                "evidence_category": "other",
                "note": "整页图片无文本层→OCR 回退（mock 空）",
            },
        )
    )
    materials.append(
        (
            "M8",
            "损坏文件.pdf",
            b"%PDF-1.4 truncated garbage",
            {
                "kind": "good",
                "upload": "created",
                "analyze": "FAILED",
                "note": "magic 通过但 pymupdf 打不开→FAILED+可重试",
            },
        )
    )
    materials.append(
        (
            "M9",
            "超大文件.txt",
            b"x" * (25 * 1024 * 1024 + 1),
            {
                "kind": "bad",
                "upload": "skipped",
                "skip_reason_contains": "大小上限",
                "note": "超过 25MB；批量接口逐文件 skipped（非整请求 413）",
            },
        )
    )
    materials.append(
        (
            "M10a",
            "伪装.txt",
            b"<html><body><script>alert(1)</script>fake txt</body></html>",
            {
                "kind": "bad",
                "upload": "skipped",
                "skip_reason_contains": "不一致",
                "note": "HTML 内容改 .txt→libmagic 判 text/html→拒绝",
            },
        )
    )
    materials.append(
        (
            "M10b",
            "伪装截图.png",
            b"#!/bin/sh\necho pwn\n",
            {
                "kind": "bad",
                "upload": "skipped",
                "skip_reason_contains": "不一致",
                "note": "脚本内容改 .png→magic 不匹配→拒绝",
            },
        )
    )
    materials.append(
        (
            "M11",
            "劳动合同（最终版）；附页 v2.pdf",
            _pdf_with_text(M11_TEXT),
            {
                "kind": "good",
                "upload": "created",
                "analyze": "SUCCESS",
                "classification": "labor_contract",
                "filename_roundtrip": True,
                "download_ascii_fallback": "v2.pdf",
                "note": "中文+空格+全角分号文件名；内容与 M1 不同避免去重",
            },
        )
    )
    materials.append(
        (
            "M12",
            "多页混合.pdf",
            _pdf_mixed(_scanned_page_png()),
            {
                "kind": "good",
                "upload": "created",
                "analyze": "SUCCESS",
                "classification": "labor_contract",
                "raw_text_include": ["第一页：劳动合同首页", "第三页：补充条款"],
                "raw_text_exclude": [M12_IMAGE_MARKER],
                "note": "3页文本/图片/文本；图片页贡献空文本（图内标记字不得出现在 raw_text）",
            },
        )
    )

    manifest: list[dict] = []
    for mid, filename, data, expect in materials:
        path = outdir / filename
        path.write_bytes(data)
        # Attack samples (kind=bad) are *supposed* to fail validation; only
        # legitimate materials must carry correct magic bytes.
        if expect["kind"] == "good" and not _magic_ok(filename, data):
            raise SystemExit(f"自校验失败：{filename} magic 不正确")
        if len(data) == 0:
            raise SystemExit(f"自校验失败：{filename} 为空")
        manifest.append(
            {
                "id": mid,
                "filename": filename,
                "size": len(data),
                "sha256": hashlib.sha256(data).hexdigest(),
                "expect": expect,
            }
        )

    (outdir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return manifest


def main() -> int:
    if len(sys.argv) != 2:
        print(__doc__, file=sys.stderr)
        return 2
    outdir = Path(sys.argv[1])
    manifest = generate(outdir)
    print(f"已生成 {len(manifest)} 份材料 → {outdir}")
    for m in manifest:
        print(f"  {m['id']:5s} {m['filename']} ({m['size']} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
